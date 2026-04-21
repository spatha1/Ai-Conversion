"""
ppt_generator.py — Build a .pptx presentation from query results and insight data.

Never writes to disk — returns raw bytes via io.BytesIO.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Optional


def build_ppt(
    title: str,
    connection_name: str,
    insights: dict,
    columns: list[str],
    rows: list[dict],
    chart_data: Optional[list[dict]] = None,
) -> bytes:
    """
    Returns raw .pptx bytes.

    Slides:
      1 — Executive Summary (title, connection, date)
      2 — Dataset Summary text
      3 — Key Insights (trends + outliers as bullets)
      4 — Data Table (first 20 rows)
      5 (optional) — Bar chart if chart_data provided
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt, Emu
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        raise RuntimeError(
            "python-pptx is not installed. Run: pip install python-pptx"
        )

    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)

    BRAND_COLOR = RGBColor(0x1E, 0x40, 0xAF)  # blue
    ACCENT_COLOR = RGBColor(0x16, 0xA3, 0x4A)  # green

    def _title_slide(prs, headline: str, subtitle: str):
        layout = prs.slide_layouts[6]  # blank
        slide = prs.slides.add_slide(layout)
        # Background rectangle
        from pptx.util import Inches
        bg = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BRAND_COLOR
        bg.line.fill.background()
        # Title
        tx = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.5))
        tf = tx.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = headline
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # Subtitle
        tx2 = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11), Inches(0.8))
        tf2 = tx2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(18)
        p2.font.color.rgb = RGBColor(0xCC, 0xDD, 0xFF)
        return slide

    def _content_slide(prs, slide_title: str, bullets: list[str], font_size: int = 18):
        layout = prs.slide_layouts[1]  # title + content
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_title
        slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = BRAND_COLOR
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = bullet
            p.font.size = Pt(font_size)
            p.level = 0
        return slide

    def _table_slide(prs, slide_title: str, cols: list[str], data_rows: list[dict]):
        layout = prs.slide_layouts[6]  # blank
        slide = prs.slides.add_slide(layout)
        tx = slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(12), Inches(0.6))
        tf = tx.text_frame
        p = tf.paragraphs[0]
        p.text = slide_title
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = BRAND_COLOR

        max_cols = min(len(cols), 8)
        max_rows = min(len(data_rows), 20)
        display_cols = cols[:max_cols]

        rows_count = max_rows + 1  # +1 header
        table = slide.shapes.add_table(rows_count, max_cols, Inches(0.3), Inches(1.0),
                                        Inches(12.7), Inches(5.8)).table
        table.columns[0].width = Inches(12.7 / max_cols)

        # Header row
        for ci, col in enumerate(display_cols):
            cell = table.cell(0, ci)
            cell.text = str(col)
            cell.text_frame.paragraphs[0].font.bold = True
            cell.text_frame.paragraphs[0].font.size = Pt(11)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BRAND_COLOR
            cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Data rows
        for ri, row in enumerate(data_rows[:max_rows]):
            for ci, col in enumerate(display_cols):
                cell = table.cell(ri + 1, ci)
                val = row.get(col)
                cell.text = "" if val is None else str(val)[:80]
                cell.text_frame.paragraphs[0].font.size = Pt(10)
                if ri % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xF0, 0xF4, 0xFF)
        return slide

    # ── Build slides ────────────────────────────────────────────
    date_str = datetime.utcnow().strftime("%Y-%m-%d")
    subtitle = f"{connection_name}  ·  {date_str}"
    _title_slide(prs, title[:120], subtitle)

    # Summary slide
    narrative = insights.get("narrative", {})
    summary_text = insights.get("summary", "")
    exec_summary = narrative.get("executive_summary", "")
    key_finding = narrative.get("key_finding", "")
    recommendation = narrative.get("recommendation", "")

    summary_bullets = [summary_text]
    if exec_summary:
        summary_bullets.append(exec_summary)
    if key_finding:
        summary_bullets.append(f"Key finding: {key_finding}")
    if recommendation:
        summary_bullets.append(f"Recommendation: {recommendation}")
    _content_slide(prs, "Dataset Summary", summary_bullets, font_size=16)

    # Key insights slide
    insight_bullets = []
    for t in insights.get("trends", [])[:4]:
        insight_bullets.append(f"Trend — {t['note']}")
    for o in insights.get("outliers", [])[:3]:
        insight_bullets.append(f"Outlier — {o['note']}")
    for n in insights.get("nulls", [])[:3]:
        insight_bullets.append(f"Data quality — '{n['column']}' has {n['null_rate']} null values")
    if not insight_bullets:
        insight_bullets = ["No significant patterns detected in this dataset."]
    _content_slide(prs, "Key Insights", insight_bullets, font_size=15)

    # Data table slide
    if rows and columns:
        _table_slide(prs, "Data Preview", columns, rows[:20])

    # Chart slide (if chart_data provided)
    if chart_data and len(chart_data) >= 2:
        try:
            from pptx.chart.data import ChartData
            from pptx.enum.chart import XL_CHART_TYPE
            layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(layout)
            tx = slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(12), Inches(0.5))
            tx.text_frame.paragraphs[0].text = "Chart"
            tx.text_frame.paragraphs[0].font.size = Pt(22)
            tx.text_frame.paragraphs[0].font.color.rgb = BRAND_COLOR

            chart_data_obj = ChartData()
            chart_data_obj.categories = [str(d.get("name", "")) for d in chart_data[:20]]
            chart_data_obj.add_series("Value", [d.get("value", 0) for d in chart_data[:20]])

            chart = slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED,
                Inches(0.5), Inches(1.0), Inches(12.0), Inches(5.8),
                chart_data_obj,
            ).chart
            chart.has_legend = False
        except Exception:
            pass  # chart generation is optional

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
